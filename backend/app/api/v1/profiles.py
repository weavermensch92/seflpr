from fastapi import APIRouter, Depends, Form, UploadFile, File, HTTPException
from typing import Optional
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.core.dependencies import get_current_user_id
from app.schemas.profile import (
    ProfileCreate, ProfileUpdate, ProfileResponse,
    FreeTextParseRequest, LinkParseRequest, FreeTextParseResponse, BulkConfirmRequest,
    IngestResponse, DashboardResponse,
)
from app.services.profile_service import ProfileService

router = APIRouter(prefix="/profiles", tags=["profiles"])


@router.get("", response_model=list[ProfileResponse])
async def list_profiles(
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    return await ProfileService(db).list_profiles(user_id)


@router.post("", response_model=ProfileResponse, status_code=201)
async def create_profile(
    body: ProfileCreate,
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    return await ProfileService(db).create_profile(user_id, body)


# ─── Unified Ingest (통합 업로드) ─────────────────────────

@router.post("/ingest", response_model=IngestResponse, status_code=201)
async def ingest(
    source_type: str = Form(..., description="file | text | link"),
    enrichment_level: str = Form("basic", description="basic | full"),
    file: Optional[UploadFile] = File(None),
    text: Optional[str] = Form(None),
    url: Optional[str] = Form(None),
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    """통합 프로필 업로드. 파일/텍스트/링크 하나로 AI 분류 + 즉시 저장."""
    if file and file.size and file.size > 100 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="파일 크기는 100MB 이하여야 합니다.")
    return await ProfileService(db).ingest(
        user_id=user_id,
        source_type=source_type,
        enrichment_level=enrichment_level,
        file=file,
        text=text,
        url=url,
    )


# ─── Dashboard ────────────────────────────────────────────

@router.get("/dashboard", response_model=DashboardResponse)
async def get_dashboard(
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    """프로필 대시보드 데이터 (타임라인, 스킬분포, 완성도, AI 강점 요약)."""
    return await ProfileService(db).get_dashboard(user_id)


@router.get("/{profile_id}", response_model=ProfileResponse)
async def get_profile(
    profile_id: str,
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    return await ProfileService(db).get_profile(profile_id, user_id)


@router.put("/{profile_id}", response_model=ProfileResponse)
async def update_profile(
    profile_id: str,
    body: ProfileUpdate,
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    return await ProfileService(db).update_profile(profile_id, user_id, body)


@router.delete("/{profile_id}", status_code=204)
async def delete_profile(
    profile_id: str,
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    await ProfileService(db).delete_profile(profile_id, user_id)


@router.post("/parse/text", response_model=FreeTextParseResponse)
async def parse_free_text(
    body: FreeTextParseRequest,
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    """AI를 사용하여 텍스트를 프로필 항목으로 파싱합니다. (포인트 보유 유저만 가능)"""
    return await ProfileService(db).parse_free_text(user_id, body)


@router.post("/parse/link", response_model=FreeTextParseResponse)
async def parse_link(
    body: LinkParseRequest,
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    """링크 URL → 텍스트 추출 → AI 파싱."""
    return await ProfileService(db).parse_url(user_id, body.url)


@router.post("/parse/file/extract")
async def extract_file_text_only(
    file: UploadFile = File(...),
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    """파일 업로드 → 텍스트 추출만 수행. (무료, AI 미사용)"""
    MAX_SIZE = 100 * 1024 * 1024  # 100MB
    content = await file.read()
    if len(content) > MAX_SIZE:
        raise HTTPException(status_code=413, detail="파일 크기는 100MB 이하여야 합니다.")
    
    extracted_text = await _extract_text(content, file.content_type or "", file.filename or "")
    return {"filename": file.filename, "text": extracted_text}


@router.post("/parse/file", response_model=FreeTextParseResponse)
async def parse_file(
    file: UploadFile = File(...),
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    """파일 업로드 → 텍스트 추출 → AI 파싱. (포인트 보유 유저만 가능)"""
    MAX_SIZE = 100 * 1024 * 1024  # 100MB
    content = await file.read()
    if len(content) > MAX_SIZE:
        raise HTTPException(status_code=413, detail="파일 크기는 100MB 이하여야 합니다.")

    extracted_text = await _extract_text(content, file.content_type or "", file.filename or "")
    return await ProfileService(db).parse_free_text(
        user_id, FreeTextParseRequest(text=extracted_text[:8000])
    )


@router.post("/parse/file/memory", response_model=ProfileResponse)
async def interpret_file_to_memory(
    file: UploadFile = File(...),
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    """파일 업로드 → 고도의 AI 해석 → 자소서용 영구 메모리로 즉시 저장. (100MB 지원, 포인트 차감)"""
    MAX_SIZE = 100 * 1024 * 1024  # 100MB
    content = await file.read()
    if len(content) > MAX_SIZE:
        raise HTTPException(status_code=413, detail="파일 크기는 100MB 이하여야 합니다.")

    extracted_text = await _extract_text(content, file.content_type or "", file.filename or "")
    if not extracted_text.strip():
        raise HTTPException(status_code=400, detail="파일에서 텍스트를 추출할 수 없습니다.")
    
    return await ProfileService(db).interpret_file_to_memory(
        user_id, file.filename or "uploaded_file", extracted_text
    )


async def _extract_text(content: bytes, content_type: str, filename: str) -> str:
    """파일 종류별 텍스트 추출."""
    import traceback
    import logging
    logger = logging.getLogger(__name__)
    
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    logger.info(f"파싱 시작: {filename} (Type: {content_type}, Ext: {ext})")

    # PDF
    if "pdf" in content_type or ext == "pdf":
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(stream=content, filetype="pdf")
            return "\n".join(page.get_text() for page in doc)
        except ImportError:
            return "[PDF 파싱: pymupdf 미설치]"

    # Excel
    if "spreadsheet" in content_type or ext in ("xlsx", "xls"):
        try:
            import openpyxl, io
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    row_str = "\t".join(str(c) if c is not None else "" for c in row)
                    if row_str.strip():
                        lines.append(row_str)
            return "\n".join(lines)
        except ImportError:
            return "[Excel 파싱: openpyxl 미설치]"

    # Word
    if "wordprocessingml" in content_type or ext == "docx":
        try:
            import docx, io
            doc = docx.Document(io.BytesIO(content))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except ImportError:
            return "[Word 파싱: python-docx 미설치]"

    # Image (OCR)
    if content_type.startswith("image/") or ext in ("jpg", "jpeg", "png", "webp"):
        try:
            from PIL import Image
            import io
            # PaddleOCR이 없을 경우 간단한 메시지 반환
            try:
                from paddleocr import PaddleOCR
                ocr = PaddleOCR(use_angle_cls=True, lang="korean", show_log=False)
                img = Image.open(io.BytesIO(content))
                import numpy as np
                result = ocr.ocr(np.array(img), cls=True)
                lines = [line[1][0] for block in result for line in block if line[1][1] > 0.5]
                return "\n".join(lines)
            except ImportError:
                logger.warning(f"PaddleOCR 미설치: {filename}")
                return "[이미지 OCR: paddleocr 미설치 - 텍스트를 직접 입력해주세요]"
        except Exception as e:
            logger.error(f"이미지 파싱 에러 ({filename}): {traceback.format_exc()}")
            return f"[이미지 처리 오류: {str(e)}]"

    logger.warning(f"지원하지 않는 파일 형식: {filename} ({content_type})")
    return "[지원하지 않는 파일 형식입니다]"


@router.post("/confirm", response_model=list[ProfileResponse], status_code=201)
async def confirm_parsed(
    body: BulkConfirmRequest,
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    """파싱 결과 확인 후 DB에 저장. 원본 파일은 이미 폐기된 상태."""
    return await ProfileService(db).bulk_confirm(user_id, body)
